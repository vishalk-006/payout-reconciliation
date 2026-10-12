"""Generate scenario folders under sample_data/scenarios/, each with a payout register and a bank statement.

    .venv/bin/python scripts/make_more_samples.py

See sample_data/scenarios/README.md (written by this script) for what each scenario tests.
"""
from __future__ import annotations

import csv
import random
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "sample_data" / "scenarios"
CHRONON_HEADER = ["Payout ID", "Employee ID", "Employee Name", "Bank Account (masked)", "IFSC", "Amount",
                  "Payout Date", "Claim IDs", "Status"]
FIRST_NAMES = ["Rahul", "Priya", "Amit", "Neha", "Karan", "Sneha", "Vikram", "Anjali", "Rohan", "Meera", "Arjun",
               "Divya", "Suresh", "Farhan", "Lakshmi", "Nikhil", "Pooja", "Ishaan", "Kavya", "Manish"]
LAST_NAMES = ["Kumar", "Shah", "Verma", "Iyer", "Mehta", "Rao", "Singh", "Desai", "Gupta", "Nair", "Das", "Pillai",
              "Babu", "Ali", "Joshi", "Kulkarni", "Reddy", "Menon", "Bose", "Chopra"]
IFSC_BANKS = ["HDFC", "ICIC", "SBIN", "UTIB", "KKBK"]


@dataclass
class Payout:
    payout_id: str
    employee_id: str
    name: str
    account: str
    ifsc: str
    amount: Decimal
    on: date
    claims: str
    status: str = "PAID"

    @property
    def last4(self) -> str:
        return self.account[-4:]


@dataclass
class Txn:
    on: date
    narration: str
    ref: str
    debit: Decimal = Decimal("0")
    credit: Decimal = Decimal("0")


def make_payouts(n: int, start: date, rng: random.Random, prefix: str = "PO-2026-") -> list[Payout]:
    payouts = []
    for i in range(n):
        first, last = rng.choice(FIRST_NAMES), rng.choice(LAST_NAMES)
        payouts.append(Payout(
            payout_id=f"{prefix}{1000 + i:04d}",
            employee_id=f"E{2000 + i}",
            name=f"{first} {last}",
            account=f"XXXXXX{rng.randint(1000, 9999)}",
            ifsc=f"{rng.choice(IFSC_BANKS)}000{rng.randint(1000, 9999)}",
            amount=Decimal(rng.randint(150, 9000) * 100 + rng.choice([0, 0, 0, 50])) / 100,
            on=start + timedelta(days=i % 10),
            claims=", ".join(f"CL-{rng.randint(10000, 99999)}" for _ in range(rng.randint(1, 3))),
        ))
    return payouts


# --------------------------------------------------------------------------- bank rows for each outcome

def paid_with_id(p: Payout, lag: int = 0) -> Txn:
    return Txn(p.on + timedelta(days=lag), f"NEFT/CHRONON/{p.payout_id}/{p.name.upper()}", f"N{p.payout_id[-4:]}",
               debit=p.amount)


def paid_by_name(p: Payout, lag: int = 1) -> Txn:
    first, last = p.name.upper().split()
    return Txn(p.on + timedelta(days=lag), f"IMPS/P2A/{first} {last[0]}/XX{p.last4}/REIMB", f"I{p.payout_id[-4:]}",
               debit=p.amount)


def paid_wrong_amount(p: Payout, delta: Decimal) -> Txn:
    return Txn(p.on, f"NEFT/{p.name.upper()}/CHRONON REIMB", f"N{p.payout_id[-4:]}", debit=p.amount + delta)


def returned(p: Payout, after_days: int = 3) -> Txn:
    return Txn(p.on + timedelta(days=after_days), f"NEFT RETURN {p.payout_id} INVALID A/C", f"R{p.payout_id[-4:]}",
               credit=p.amount)


def unrelated(on: date, rng: random.Random) -> Txn:
    kind = rng.choice(["charges", "salary", "vendor", "interest", "gst"])
    if kind == "charges":
        return Txn(on, "CONSOLIDATED CHARGES FOR A/C", "", debit=Decimal(rng.choice([59, 118, 236])))
    if kind == "salary":
        return Txn(on, "NEFT SALARY ADVANCE OPS TEAM", f"S{rng.randint(100, 999)}", debit=Decimal(rng.randint(20, 90) * 1000))
    if kind == "vendor":
        return Txn(on, f"RTGS ACME SUPPLIES PVT LTD INV {rng.randint(1000, 9999)}", f"V{rng.randint(100, 999)}",
                   debit=Decimal(rng.randint(10, 200) * 500))
    if kind == "gst":
        return Txn(on, "GST PAYMENT CHALLAN", f"G{rng.randint(100, 999)}", debit=Decimal(rng.randint(5, 50) * 1000))
    return Txn(on, "INT.PD FOR THE QUARTER", "", credit=Decimal(rng.randint(100, 2000)) + Decimal("0.35"))


def with_balances(txns: list[Txn], opening: Decimal) -> list[tuple[Txn, Decimal]]:
    rows, balance = [], opening
    for t in sorted(txns, key=lambda t: t.on):
        balance = balance - t.debit + t.credit
        rows.append((t, balance))
    return rows


def amount_text(v: Decimal) -> str | None:
    return f"{v:,.2f}" if v else None


# --------------------------------------------------------------------------- file writers (one per bank format)

def write_register(path: Path, payouts: list[Payout], header: list[str] = CHRONON_HEADER,
                   reorder: Callable[[list], list] = lambda r: r) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Payouts"
    ws.append(reorder(header))
    for p in payouts:
        ws.append(reorder([p.payout_id, p.employee_id, p.name, p.account, p.ifsc, float(p.amount),
                           p.on.strftime("%d-%m-%Y"), p.claims, p.status]))
    wb.save(path)


def write_seeded_format(path: Path, txns: list[Txn], opening: Decimal, extra=None) -> None:
    """The layout the service already knows (no LLM call needed)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["ABC BANK LTD - ACCOUNT STATEMENT"])
    ws.append(["Account No: 50100XXXXXX123", None, "Customer: CHRONON TECHNOLOGIES PVT LTD"])
    ws.append([])
    header = ["Txn Date", "Value Date", "Narration", "Cheque/Ref No", "Debit", "Credit", "Balance"]
    ws.append(header)
    ws.append([None, None, "Opening Balance", None, None, None, f"{opening:,.2f}"])
    rows = with_balances(txns, opening)
    for t, bal in rows:
        d = t.on.strftime("%d-%m-%Y")
        ws.append([d, d, t.narration, t.ref, amount_text(t.debit), amount_text(t.credit), f"{bal:,.2f}"])
    ws.append([None, None, "Closing Balance", None, None, None, f"{rows[-1][1]:,.2f}"])
    if extra:
        extra(wb, ws, rows)
    wb.save(path)


def write_hdfc_style(path: Path, txns: list[Txn], opening: Decimal) -> None:
    """Different names and order, 2-digit years, numeric cells, a statement summary footer."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Account Statement"
    ws.append(["HDFC BANK Ltd.", None, None, "Page No .: 1"])
    ws.append(["M/S. CHRONON TECHNOLOGIES PVT LTD", None, None, "Account No : 50200012345678"])
    ws.append(["Statement From : 01/03/26 To : 15/03/26"])
    ws.append([])
    ws.append(["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"])
    ws.append(["*" * 20])
    rows = with_balances(txns, opening)
    for t, bal in rows:
        d = t.on.strftime("%d/%m/%y")
        ws.append([d, t.narration, t.ref, d, float(t.debit) or None, float(t.credit) or None, float(bal)])
    ws.append([])
    ws.append(["STATEMENT SUMMARY :-"])
    ws.append(["Opening Balance", "Dr Count", "Cr Count", "Debits", "Credits", "Closing Bal"])
    debits, credits = sum(t.debit for t in txns), sum(t.credit for t in txns)
    ws.append([float(opening), sum(1 for t in txns if t.debit), sum(1 for t in txns if t.credit), float(debits),
               float(credits), float(rows[-1][1])])
    ws.append(["Generated On: 16-Mar-2026 10:12:44  This is a computer generated statement."])
    wb.save(path)


def write_icici_style(path: Path, txns: list[Txn], opening: Decimal) -> None:
    """Newest first, timestamps in the date column, one amount column plus a Cr/Dr column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "OpTransactionHistory"
    ws.append(["ICICI Bank", None, "Detailed Statement"])
    ws.append(["Transactions List - CHRONON TECHNOLOGIES (INR) - 001105012345"])
    ws.append([])
    ws.append(["S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks",
               "Withdrawal/Deposit", "Cr/Dr", "Balance (INR )"])
    rows = list(reversed(with_balances(txns, opening)))
    for i, (t, bal) in enumerate(rows, start=1):
        stamp = datetime.combine(t.on, datetime.min.time()) + timedelta(hours=9, minutes=i % 60)
        amount = t.debit or t.credit
        ws.append([i, t.on.strftime("%d/%m/%Y"), stamp.strftime("%d/%m/%Y %I:%M:%S %p"), t.ref or "-", t.narration,
                   f"{amount:,.2f}", "DR" if t.debit else "CR", f"{bal:,.2f}"])
    ws.append([])
    ws.append(["Legends Used in Account Statement"])
    wb.save(path)


def write_signed_csv(path: Path, txns: list[Txn], opening: Decimal) -> None:
    """CSV export: a single signed amount (negative = money out) and month-name dates."""
    rows = with_balances(txns, opening)
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Account Activity Export"])
        w.writerow([])
        w.writerow(["Posting Date", "Description", "Reference", "Amount", "Running Balance"])
        for t, bal in rows:
            signed = t.credit - t.debit
            w.writerow([t.on.strftime("%d-%b-%Y"), t.narration, t.ref, f"{signed:.2f}", f"{bal:.2f}"])


def write_no_balance(path: Path, txns: list[Txn]) -> None:
    """No balance column: nothing in the file can prove debit and credit are the right way round."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Txns"
    ws.append(["Date", "Particulars", "Ref", "Debit (INR)", "Credit (INR)"])
    for t in sorted(txns, key=lambda t: t.on):
        ws.append([t.on, t.narration, t.ref, float(t.debit) or None, float(t.credit) or None])
    wb.save(path)


# --------------------------------------------------------------------------- scenarios

def mixed_outcomes(payouts: list[Payout], rng: random.Random) -> list[Txn]:
    """A realistic week: most paid with ID, some by name only, plus every kind of exception."""
    txns: list[Txn] = []
    for i, p in enumerate(payouts):
        slot = i % 10
        if slot in (0, 1, 2, 3, 4):
            txns.append(paid_with_id(p, lag=i % 2))
        elif slot in (5, 6):
            txns.append(paid_by_name(p))
        elif slot == 7:
            txns.append(paid_wrong_amount(p, Decimal("50.00")))
        elif slot == 8:
            txns += [paid_with_id(p), returned(p)]
        else:
            pass  # missing in bank
    for d in range(10):
        if rng.random() < 0.6:
            txns.append(unrelated(payouts[0].on + timedelta(days=d), rng))
    return txns


def scenario_all_matched(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(12, date(2026, 3, 2), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    write_seeded_format(folder / "bank_statement.xlsx", [paid_with_id(p) for p in payouts], Decimal("500000.00"))
    return "Known layout. Every payout paid once with its Payout ID: everything matches, nothing needs review."


def scenario_hdfc(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(20, date(2026, 3, 2), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    write_hdfc_style(folder / "bank_statement.xlsx", mixed_outcomes(payouts, rng), Decimal("750000.00"))
    return ("NEW bank layout (needs an LLM key): 'Withdrawal Amt.'/'Deposit Amt.', dd/mm/yy dates, numeric cells, "
            "statement-summary footer. Mixed outcomes: ID matches, name/account matches, amount mismatch, "
            "returned payout, missing payouts, unrelated debits.")


def scenario_icici(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(15, date(2026, 3, 9), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    write_icici_style(folder / "bank_statement.xlsx", mixed_outcomes(payouts, rng), Decimal("320000.00"))
    return ("NEW bank layout (needs an LLM key): newest-first, timestamps in the transaction date, one amount column "
            "plus a Cr/Dr column, two date columns. Mixed outcomes.")


def scenario_signed_csv(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(10, date(2026, 3, 16), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    write_signed_csv(folder / "bank_statement.csv", mixed_outcomes(payouts, rng), Decimal("100000.00"))
    return ("NEW bank layout (needs an LLM key), CSV: single signed Amount column (negative = debit), "
            "dd-Mon-yyyy dates. Mixed outcomes.")


def scenario_broken_statement(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(12, date(2026, 3, 2), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    txns = [paid_with_id(p) for p in payouts]

    def damage(wb, ws, rows):
        data_start = 6  # header on row 4, opening balance on row 5
        ws.delete_rows(data_start + 4)  # a transaction lost from the export: the running balance breaks
        dup = [c.value for c in ws[data_start + 1]]
        ws.insert_rows(data_start + 2)  # the same line exported twice
        for col, value in enumerate(dup, start=1):
            ws.cell(data_start + 2, col, value)
        ws.cell(data_start + 7, 1, "31/02/2026")  # an impossible date
        ws.cell(data_start + 9, 5, "1,2O0.00")  # letter O instead of zero
        wb.create_sheet("Page 2").append(["16-03-2026", "16-03-2026", "NEFT/CHRONON/CONTINUED", "X1", "999.00", None,
                                          "0.00"])
    write_seeded_format(folder / "bank_statement.xlsx", txns, Decimal("400000.00"), extra=damage)
    return ("Known layout, damaged export: a missing row (balance break), a duplicated row, an impossible date, "
            "an amount typed with a letter O, and a second sheet that is not read. Expect P1 integrity items, "
            "unreadable rows listed, and totals that still tie.")


def scenario_no_balance(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(8, date(2026, 3, 2), rng)
    write_register(folder / "payout_register.xlsx", payouts)
    write_no_balance(folder / "bank_statement.xlsx", [paid_with_id(p) for p in payouts])
    return ("NEW bank layout (needs an LLM key) with NO balance column: the layout cannot be proven, so it is "
            "saved as 'pending review' and flagged until approved with scripts/approve_layout.py.")


def scenario_renamed_register(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(10, date(2026, 3, 2), rng)
    header = ["Beneficiary", "Payout Ref", "Net Amount", "Paid On", "Payout Status", "Emp Code", "A/c No",
              "IFSC Code", "Expense Claims"]
    write_register(folder / "payout_register.xlsx", payouts, header=header,
                   reorder=lambda r: r if r is header else [r[2], r[0], r[5], r[6], r[8], r[1], r[3], r[4], r[7]])
    write_seeded_format(folder / "bank_statement.xlsx", [paid_with_id(p) for p in payouts], Decimal("200000.00"))
    return ("Payout register exported with renamed and reordered columns (needs an LLM key). The bank layout is "
            "known. A new register layout is always held for review.")


def scenario_large(folder: Path, rng: random.Random) -> str:
    payouts = make_payouts(3000, date(2026, 3, 1), rng)
    for i, p in enumerate(payouts):
        p.on = date(2026, 3, 1) + timedelta(days=i % 28)
        if i % 97 == 0:
            p.status = "FAILED"
        elif i % 89 == 0:
            p.status = "PENDING"
    txns = []
    for p in payouts:
        if p.status != "PAID":
            continue
        r = rng.random()
        if r < 0.80:
            txns.append(paid_with_id(p, lag=rng.choice([0, 0, 1, 2])))
        elif r < 0.93:
            txns.append(paid_by_name(p, lag=rng.choice([0, 1, 2])))
        elif r < 0.96:
            txns.append(paid_wrong_amount(p, Decimal(rng.choice([-100, 50, 10]))))
        elif r < 0.98:
            txns += [paid_with_id(p), returned(p, after_days=rng.randint(1, 5))]
    txns += [unrelated(date(2026, 3, 1) + timedelta(days=rng.randint(0, 27)), rng) for _ in range(120)]
    write_register(folder / "payout_register.xlsx", payouts)
    write_seeded_format(folder / "bank_statement.xlsx", txns, Decimal("25000000.00"))
    return "Known layout, volume test: 3,000 payouts over a month with a realistic mix of outcomes."


SCENARIOS: dict[str, Callable[[Path, random.Random], str]] = {
    "01_all_matched": scenario_all_matched,
    "02_hdfc_style_new_layout": scenario_hdfc,
    "03_icici_style_newest_first": scenario_icici,
    "04_signed_amount_csv": scenario_signed_csv,
    "05_broken_statement": scenario_broken_statement,
    "06_no_balance_column": scenario_no_balance,
    "07_renamed_register": scenario_renamed_register,
    "08_large_volume": scenario_large,
}


def main() -> None:
    lines = ["# Sample scenarios", "", "Generated by `scripts/make_more_samples.py`. Upload the two files of a folder "
             "with `POST /api/files`, then reconcile.", "",
             "Scenarios marked *needs an LLM key* use a bank or register layout the service has not seen; set "
             "`GEMINI_API_KEY` (or `LLM_PROVIDER=claude` + `ANTHROPIC_API_KEY`) first, otherwise the API answers 503.",
             ""]
    for name, build in SCENARIOS.items():
        folder = OUT / name
        folder.mkdir(parents=True, exist_ok=True)
        description = build(folder, random.Random(name))
        lines.append(f"- **{name}**: {description}")
        print(f"{name}: {', '.join(sorted(p.name for p in folder.iterdir()))}")
    (OUT / "README.md").write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    main()
