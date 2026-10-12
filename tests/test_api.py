import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import make_sample_data as sample


@pytest.fixture
def client():
    from app.main import app
    return TestClient(app)  # storage is isolated per test by conftest


@pytest.fixture
def sample_files(tmp_path):
    p, b = tmp_path / "payout_register.xlsx", tmp_path / "bank_statement.xlsx"
    sample.make_payout_register(p)
    sample.make_bank_statement(b)
    return p, b


def upload(client, path):
    with open(path, "rb") as f:
        return client.post("/api/files", files={"file": (path.name, f)})


def test_upload_then_reconcile_returns_excel_report(client, sample_files):
    p, b = sample_files
    up_p, up_b = upload(client, p), upload(client, b)
    assert up_p.status_code == 201 and up_b.status_code == 201
    assert up_p.json()["detected_type"] == "payout_register"
    assert up_b.json()["detected_type"] == "bank_statement" and up_b.json()["layout_id"]

    resp = client.post("/api/reconciliations",
                       json={"payout_file_id": up_p.json()["file_id"], "bank_file_id": up_b.json()["file_id"]})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.headers["x-totals-tie"] == "true"
    assert resp.headers["x-bank-balance-check"] == "VERIFIED"
    assert resp.headers["x-layouts-verified"] == "true"

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Action List (Needs Review)", "Summary", "Matched", "All Payouts", "All Bank Rows",
                             "Data Quality"]
    summary_text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "DOES NOT TIE" not in summary_text and summary_text.count("✔ ties") == 3
    actions = " ".join(str(c.value) for row in wb["Action List (Needs Review)"].iter_rows() for c in row if c.value)
    for heading in ("Pay these employees again", "Recover money", "Find out what this money is", "Fix the files",
                    "Ask the bank", "Confirm these matches", "check again on the next statement"):
        assert heading in actions
    assert "The bank sent ₹3,100.00" in actions and "came back" in actions  # the returned payout, in plain words

    again = client.get(f"/api/reports/{resp.headers['x-report-id']}")
    assert again.status_code == 200 and again.content == resp.content


def test_same_file_uploaded_twice_reuses_id(client, sample_files):
    a, b = upload(client, sample_files[0]), upload(client, sample_files[0])
    assert a.json()["file_id"] == b.json()["file_id"] and b.json()["already_uploaded"] is True


def test_rejects_non_excel_and_corrupt_files(client):
    assert client.post("/api/files", files={"file": ("x.pdf", b"%PDF")}).status_code == 415
    assert client.post("/api/files", files={"file": ("x.xlsx", b"not a zip")}).status_code == 422


def test_unknown_ids_and_wrong_files(client, sample_files):
    r = client.post("/api/reconciliations", json={"payout_file_id": "0" * 32, "bank_file_id": "1" * 32})
    assert r.status_code == 404
    p = upload(client, sample_files[0]).json()["file_id"]
    b = upload(client, sample_files[1]).json()["file_id"]
    swapped = client.post("/api/reconciliations", json={"payout_file_id": b, "bank_file_id": p})
    assert swapped.status_code == 503  # neither layout is known for the swapped kind, and no LLM is configured
    same = client.post("/api/reconciliations", json={"payout_file_id": p, "bank_file_id": p})
    assert same.status_code == 422


def test_path_traversal_ids_are_rejected(client):
    assert client.get("/api/reports/..%2F..%2Fetc%2Fpasswd").status_code == 404


def new_format_statement(path):
    from tests.conftest import write_xlsx
    return write_xlsx(path, [["Posted", "Details", "Out", "In", "Bal"],
                             ["03-02-2026", "NEFT/CHRONON/PO-2026-0101/RAHUL KUMAR", "4,250.00", None, "5,750.00"]])


def test_new_bank_format_without_llm_returns_503(client, sample_files, tmp_path):
    p = upload(client, sample_files[0]).json()["file_id"]
    up = upload(client, new_format_statement(tmp_path / "other_bank.xlsx")).json()
    assert up["detected_type"] is None  # unknown layout: identified during reconciliation
    r = client.post("/api/reconciliations", json={"payout_file_id": p, "bank_file_id": up["file_id"]})
    assert r.status_code == 503 and "not seen before" in r.json()["detail"]


def test_new_bank_format_with_llm_provider(client, sample_files, tmp_path):
    from app.layout.providers import register
    from tests.conftest import FakeDetector, answer
    llm = FakeDetector(answer("Sheet", 1, {"txn_date": "Posted", "narration": "Details", "debit": "Out",
                                           "credit": "In", "balance": "Bal"}))
    register("fake-api")(lambda: llm)
    p = upload(client, sample_files[0]).json()["file_id"]
    b = upload(client, new_format_statement(tmp_path / "other_bank.xlsx")).json()["file_id"]
    r = client.post("/api/reconciliations", json={"payout_file_id": p, "bank_file_id": b, "llm_provider": "fake-api"})
    assert r.status_code == 200, r.text
    summary = load_workbook(io.BytesIO(r.content))["Summary"]
    text = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)
    assert "NEW layout identified by fake:test-model" in text
    # Same format again: cached, no second LLM call.
    b2 = upload(client, new_format_statement(tmp_path / "other_bank_2.xlsx")).json()
    assert b2["detected_type"] == "bank_statement"
    assert len(llm.requests) == 1


def test_unreadable_payout_rows_show_their_payout_id(client, sample_files):
    p = upload(client, sample_files[0]).json()["file_id"]
    b = upload(client, sample_files[1]).json()["file_id"]
    r = client.post("/api/reconciliations", json={"payout_file_id": p, "bank_file_id": b})
    sheet = load_workbook(io.BytesIO(r.content))["All Payouts"]
    unreadable = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[11] == "Could not read row"]
    assert {row[1] for row in unreadable} == {"PO-2026-0117", "PO-2026-0114"}
