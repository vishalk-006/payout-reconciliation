"""Excel report. Every row points back to its source file and Excel row number so
finance can verify any result in the original file within seconds."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, NamedTuple
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .actions import Action, ActionItem, build_actions, headline
from .matching import BANK_REVIEW, PAYOUT_REVIEW, BankBucket, PayoutBucket, ReconResult, day, money
from .layout.models import SEEDED
from .models import BalanceStatus, BankParseResult, Direction, LayoutInfo, PayoutParseResult

ZERO = Decimal("0")
MONEY_FMT = "#,##0.00"
DATE_FMT = "dd-mmm-yyyy"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
BOLD = Font(bold=True)
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
ACTION_FILLS = {
    Action.PAY_AGAIN: "F8CBAD", Action.RECOVER: "F8CBAD", Action.INVESTIGATE: "F8CBAD", Action.DO_NOT_REPAY: "F8CBAD",
    Action.FIX_FILE: "FCE4D6", Action.ASK_BANK: "FFE699", Action.PAY_DIFFERENCE: "FFE699",
    Action.UPDATE_STATUS: "DDEBF7", Action.CONFIRM: "DDEBF7", Action.WAIT: "EDEDED",
}
NOT_ACTIONABLE_NOW = {Action.WAIT}


@dataclass
class ReportInput:
    payouts: PayoutParseResult
    bank: BankParseResult
    result: ReconResult
    generated_at: datetime


def build_report(data: ReportInput, out_path: Path) -> dict[str, Any]:
    actions = build_actions(data.payouts, data.bank, data.result)
    wb = Workbook()
    wb.active.title = ACTION_SHEET
    _write_action_list(wb.active, data, actions)
    summary = _write_summary(wb.create_sheet("Summary"), data, actions)
    _write_matched(wb.create_sheet("Matched"), data)
    _write_all_payouts(wb.create_sheet("All Payouts"), data)
    _write_all_bank(wb.create_sheet("All Bank Rows"), data)
    _write_data_quality(wb.create_sheet("Data Quality"), data)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return summary


class OutcomeRow(NamedTuple):
    outcome: str
    count: int
    amount: Decimal
    needs_review: str


@dataclass
class TieOut:
    """Outcome buckets for one side of the reconciliation, and the input file's own totals to prove against."""
    rows: list[OutcomeRow]
    file_count: int
    file_amount: Decimal

    @property
    def count(self) -> int:
        return sum(r.count for r in self.rows)

    @property
    def amount(self) -> Decimal:
        return sum((r.amount for r in self.rows), ZERO)

    @property
    def ties(self) -> bool:
        return self.count == self.file_count and self.amount == self.file_amount

    def as_dict(self) -> dict[str, Any]:
        return {"rows": self.file_count, "amount": str(self.file_amount), "ties": self.ties}


UNREADABLE = "Could not read row (see Data Quality)"
CREDIT_BUCKETS = (BankBucket.RETURN_LINKED, BankBucket.RETURN_UNLINKED, BankBucket.OTHER_CREDIT)
DEBIT_BUCKETS = tuple(b for b in BankBucket if b not in CREDIT_BUCKETS)
BANK_BUCKETS_NEEDING_REVIEW = {*BANK_REVIEW, BankBucket.MATCHED_REVIEW}


def _yes_no(flag: bool) -> str:
    return "Yes" if flag else "No"


def payout_tie_out(data: ReportInput) -> TieOut:
    pr, res = data.payouts, data.result
    rows = []
    for bucket in PayoutBucket:
        items = [x for x in res.payouts if x.bucket is bucket]
        rows.append(OutcomeRow(bucket.value, len(items), sum((x.payout.amount for x in items), ZERO),
                               _yes_no(bucket in PAYOUT_REVIEW)))
    rejected_amount = sum((x.amount for x in pr.rejected if x.amount), ZERO)
    rows.append(OutcomeRow(UNREADABLE, len(pr.rejected), rejected_amount, "Yes"))
    return TieOut(rows, len(pr.rows) + len(pr.rejected), sum((p.amount for p in pr.rows), ZERO) + rejected_amount)


def bank_tie_out(data: ReportInput, direction: Direction) -> TieOut:
    br, res = data.bank, data.result
    buckets = DEBIT_BUCKETS if direction is Direction.DEBIT else CREDIT_BUCKETS
    rows = []
    for bucket in buckets:
        items = [x for x in res.txns if x.bucket is bucket and x.txn.direction is direction]
        rows.append(OutcomeRow(bucket.value, len(items), sum((x.txn.amount for x in items), ZERO),
                               _yes_no(bucket in BANK_BUCKETS_NEEDING_REVIEW)))
    rejected = [x for x in br.rejected if x.direction is direction]
    rejected_amount = sum((x.amount for x in rejected), ZERO)
    rows.append(OutcomeRow(UNREADABLE, len(rejected), rejected_amount, "Yes"))
    parsed = [t for t in br.rows if t.direction is direction]
    return TieOut(rows, len(parsed) + len(rejected), sum((t.amount for t in parsed), ZERO) + rejected_amount)


def _write_summary(ws: Worksheet, data: ReportInput, actions: list[ActionItem]) -> dict[str, Any]:
    pr, br, res = data.payouts, data.bank, data.result
    ws["A1"] = "Summary – where every payout and every bank row ended up, tied to the file totals"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {data.generated_at:%d-%b-%Y %H:%M}"
    period = f"{day(res.period[0])} to {day(res.period[1])}" if res.period else "—"
    r = _kv(ws, 4, [
        ("Payout register", f"{pr.file_name} (sheet '{pr.sheet}', header on row {pr.header_row}, dates {pr.date_format or 'native'})"),
        ("Bank statement", f"{br.file_name} (sheet '{br.sheet}', header on row {br.header_row}, "
                           f"dates {br.date_format or 'native'}, amounts '{br.amount_style}')"),
        ("Statement period", period),
    ])
    for label, info in (("Payout register layout", pr.layout), ("Bank statement layout", br.layout)):
        r = _status_line(ws, r, label, describe_layout(info), ok=info is None or info.verified)
    bc = br.balance_check
    r = _status_line(ws, r, "Bank balance check", f"{bc.status.value}: {bc.message}",
                     ok=bc.status is BalanceStatus.VERIFIED)

    payouts = payout_tie_out(data)
    r = _write_tie_out(ws, r + 1, "File A – Payout Register: where every payout ended up", "Payouts", payouts,
                       "File A total")
    if unknown := sum(1 for x in pr.rejected if x.amount is None):
        r = _note(ws, r, f"Note: {unknown} unreadable row(s) have no readable amount and are counted but not summed.")
    if pr.control_total is not None:
        ok = pr.control_total == payouts.file_amount
        ws.cell(r, 1, "Register's own Total row")
        ws.cell(r, 3, float(pr.control_total)).number_format = MONEY_FMT
        ws.cell(r, 4, "✔ agrees" if ok else "✘ DOES NOT AGREE").fill = OK_FILL if ok else BAD_FILL
        r += 1

    debits = bank_tie_out(data, Direction.DEBIT)
    r = _write_tie_out(ws, r + 1, "File B – Bank Statement debits: where every debit ended up", "Rows", debits,
                       "File B total debits")
    credits = bank_tie_out(data, Direction.CREDIT)
    r = _write_tie_out(ws, r + 1, "File B – Bank Statement credits", "Rows", credits, "File B total credits")
    if no_direction := sum(1 for x in br.rejected if x.direction is None):
        r = _note(ws, r, f"Note: {no_direction} unreadable bank row(s) have no readable amount/direction "
                         f"(see Data Quality); they are not in the debit or credit totals.")

    if bc.opening is not None and bc.closing is not None:
        r = _write_balance_proof(ws, r + 1, bc.opening, debits.file_amount, credits.file_amount, bc.closing)

    r = _section(ws, r + 1, f"Needs a person – by action (details on '{ACTION_SHEET}')")
    groups = [OutcomeRow(a.value, len(g), sum((i.amount for i in g if i.amount), ZERO), "Yes")
              for a in Action if (g := [i for i in actions if i.action is a])]
    r = _table(ws, r, ["Action", "Items", "Amount involved", ""], groups or [["Nothing needs a person", 0, None, ""]],
               money_cols={3})
    _widths(ws, [44, 70, 18, 14, 14])
    return {
        "payouts": payouts.as_dict(),
        "bank_debits": debits.as_dict(),
        "bank_credits": credits.as_dict(),
        "balance_check": bc.status.value,
        "layouts_verified": all(i is None or i.verified for i in (pr.layout, br.layout)),
        "needs_review": sum(1 for i in actions if i.action not in NOT_ACTIONABLE_NOW),
        "outcomes": {b.value: sum(1 for x in res.payouts if x.bucket is b) for b in PayoutBucket},
    }


def describe_layout(info: LayoutInfo | None) -> str:
    if info is None:
        return "—"
    if info.verified:
        if info.source == SEEDED:
            how = "Reviewed layout, shipped with the service"
        elif info.llm_called:
            how = (f"NEW layout identified by {info.source} in {info.attempts} attempt(s); proven on this file "
                   f"and saved for next time")
        else:
            how = f"Known layout, identified earlier by {info.source} and proven on a full file"
    else:
        origin = f"identified now by {info.source}" if info.llm_called else f"from cache ({info.source})"
        how = f"Layout {origin} is NOT VERIFIED: {' '.join(info.problems)}"
    return f"{how}. [{info.id}]" + (f" Notes: {info.notes}" if info.notes else "")


def _write_tie_out(ws: Worksheet, r: int, title: str, unit: str, tie: TieOut, file_label: str) -> int:
    r = _section(ws, r, title)
    r = _table(ws, r, ["Outcome", unit, "Amount", "Needs review"], tie.rows, money_cols={3})
    for i, (label, count, amount) in enumerate((("Total of outcomes", tie.count, tie.amount),
                                                (file_label, tie.file_count, tie.file_amount))):
        ws.cell(r + i, 1, label).font = BOLD
        ws.cell(r + i, 2, count).font = BOLD
        cell = ws.cell(r + i, 3, float(amount))
        cell.number_format, cell.font = MONEY_FMT, BOLD
    ws.cell(r + 1, 4, "✔ ties" if tie.ties else "✘ DOES NOT TIE").fill = OK_FILL if tie.ties else BAD_FILL
    return r + 2


def _write_balance_proof(ws: Worksheet, r: int, opening: Decimal, debits: Decimal, credits: Decimal,
                         closing: Decimal) -> int:
    r = _section(ws, r, "Bank balance proof (uses the statement's own opening and closing balance)")
    computed = opening - debits + credits
    r = _table(ws, r, ["Item", "Amount"], [
        ["Opening balance", opening], ["− Total debits", debits], ["+ Total credits", credits],
        ["= Computed closing balance", computed], ["Closing balance on statement", closing],
    ], money_cols={2})
    ok = computed == closing
    ws.cell(r, 1, "Result").font = BOLD
    ws.cell(r, 2, "✔ agrees" if ok else f"✘ off by {money(closing - computed)}").fill = OK_FILL if ok else BAD_FILL
    return r + 1


def _status_line(ws: Worksheet, r: int, label: str, text: str, ok: bool) -> int:
    r = _kv(ws, r, [(label, text)])
    ws.cell(r - 1, 2).fill = OK_FILL if ok else BAD_FILL
    return r


def _note(ws: Worksheet, r: int, text: str) -> int:
    ws.cell(r, 1, text)
    return r + 1


def _write_action_list(ws: Worksheet, data: ReportInput, actions: list[ActionItem]) -> None:
    res = data.result
    ok = sum(1 for x in res.payouts if x.bucket in (PayoutBucket.MATCHED, PayoutBucket.MATCHED_BATCH,
                                                    PayoutBucket.RETURN_CONSISTENT, PayoutBucket.NO_DEBIT_EXPECTED))
    period = f"{day(res.period[0])} to {day(res.period[1])}" if res.period else "—"
    ws["A1"] = "What to do"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Bank statement {period} · {data.bank.file_name} vs {data.payouts.file_name} · generated {data.generated_at:%d-%b-%Y %H:%M}"
    r = 4
    for i, line in enumerate(headline(actions, len(res.payouts) + len(data.payouts.rejected), ok)):
        ws.cell(r, 1, line).font = BOLD if i == 0 else Font()
        r += 1

    for number, action in enumerate((a for a in Action if any(i.action is a for i in actions)), start=1):
        group = [i for i in actions if i.action is action]
        total = sum((i.amount for i in group if i.amount), ZERO)
        r += 1
        band = f"{number}. {action.value}  ({len(group)}" + (f" · {money(total)})" if total else ")")
        for col in range(1, len(ACTION_HEADERS) + 1):
            ws.cell(r, col).fill = PatternFill("solid", fgColor=ACTION_FILLS[action])
        ws.cell(r, 1, band).font = Font(bold=True, size=12)
        r = _table(ws, r + 1, ACTION_HEADERS,
                   [[i.employee or "—", i.payout_id or "—", i.amount, i.what_happened, i.next_step, i.evidence]
                    for i in group], money_cols={3}, wrap_cols={4, 5, 6})

    r += 1
    for line in ("How to read the other sheets:",
                 "• Summary – counts and amounts per category; proves every payout and bank row is counted exactly once (all lines should say ✔).",
                 "• Matched – payouts that are fine, next to the bank row that paid them.",
                 "• All Payouts / All Bank Rows – every row of both files and what happened to it, with its row number.",
                 "• Data Quality – rows that were skipped or could not be read, and why."):
        r = _note(ws, r, line)
    _widths(ws, [24, 16, 14, 62, 55, 44])
    ws.freeze_panes = "A4"


ACTION_HEADERS = ["Employee", "Payout ID", "Amount", "What happened", "What to do", "Where to look"]
ACTION_SHEET = "Action List (Needs Review)"


def _write_matched(ws: Worksheet, data: ReportInput) -> None:
    rows = []
    for x in data.result.payouts:
        if x.bucket not in (PayoutBucket.MATCHED, PayoutBucket.MATCHED_BATCH):
            continue
        p, t = x.payout, x.txns[0]
        rows.append([p.payout_id, p.employee_id, p.employee_name, p.account_masked, p.amount, p.payout_date,
                     ", ".join(p.claim_ids), p.source_row, t.source_row, t.txn_date, t.amount, t.narration,
                     t.reference, (t.txn_date - p.payout_date).days, x.rule, x.confidence, x.reason])
    _table(ws, 1, ["Payout ID", "Employee ID", "Employee", "Account", "Payout Amount", "Payout Date", "Claim IDs",
                   "Register Row", "Bank Row", "Bank Date", "Bank Amount", "Bank Narration", "Bank Ref",
                   "Days (Bank − Payout)", "Rule", "Confidence", "Why"], rows,
           money_cols={5, 11}, date_cols={6, 10}, wrap_cols={17})
    _finish(ws, [14, 12, 22, 16, 14, 13, 22, 10, 9, 13, 14, 40, 16, 10, 26, 11, 70])


def _write_all_payouts(ws: Worksheet, data: ReportInput) -> None:
    rows = []
    for x in sorted(data.result.payouts, key=lambda x: x.payout.source_row):
        p = x.payout
        rows.append([p.source_row, p.payout_id, p.employee_id, p.employee_name, p.account_masked, p.ifsc, p.amount,
                     p.payout_date, ", ".join(p.claim_ids), len(p.claim_ids), p.raw_status, x.bucket.value,
                     "Yes" if x.needs_review else "No",
                     ", ".join(str(t.source_row) for t in x.txns) or None, x.rule, x.confidence, x.reason,
                     "; ".join(p.warnings)])
    for rj in data.payouts.rejected:
        rows.append([rj.source_row, rj.record_id, "", "", "", "", rj.amount, None, "", None,
                     "", "Could not read row", "Yes", None, "Parsing", "", rj.reason, ""])
    rows.sort(key=lambda r: r[0])
    _table(ws, 1, ["Register Row", "Payout ID", "Employee ID", "Employee", "Account", "IFSC", "Amount", "Payout Date",
                   "Claim IDs", "# Claims", "Register Status", "Outcome", "Needs Review", "Bank Row(s)", "Rule",
                   "Confidence", "Why", "Data Warnings"], rows,
           money_cols={7}, date_cols={8}, wrap_cols={17})
    _finish(ws, [10, 14, 12, 22, 16, 13, 14, 13, 22, 8, 12, 34, 8, 12, 26, 11, 70, 30])


def _write_all_bank(ws: Worksheet, data: ReportInput) -> None:
    rows = []
    for x in data.result.txns:
        t = x.txn
        rows.append([t.source_row, t.txn_date, t.value_date, t.narration, t.reference,
                     t.debit or None, t.credit or None, t.balance, x.category.value, x.bucket.value,
                     "Yes" if x.needs_review or x.bucket is BankBucket.MATCHED_REVIEW else "No",
                     ", ".join(x.payout_ids), x.rule, x.reason, "; ".join(t.warnings)])
    for rj in data.bank.rejected:
        rows.append([rj.source_row, None, None, "; ".join(f"{k}={v}" for k, v in rj.raw.items() if v), "",
                     rj.amount if rj.direction is Direction.DEBIT else None,
                     rj.amount if rj.direction is Direction.CREDIT else None, None, "", "Could not read row", "Yes",
                     "", "Parsing", rj.reason, ""])
    rows.sort(key=lambda r: r[0])
    _table(ws, 1, ["Bank Row", "Txn Date", "Value Date", "Narration", "Reference", "Debit", "Credit", "Balance",
                   "Category", "Outcome", "Needs Review", "Payout ID(s)", "Rule", "Why", "Data Warnings"], rows,
           money_cols={6, 7, 8}, date_cols={2, 3}, wrap_cols={4, 14})
    _finish(ws, [9, 13, 13, 44, 16, 14, 14, 15, 22, 32, 8, 18, 26, 60, 30])


def _write_data_quality(ws: Worksheet, data: ReportInput) -> None:
    r = 1
    r = _section(ws, r, "Row accounting – every row of each file is accounted for exactly once")
    rows = []
    for name, acc in ((data.payouts.file_name, data.payouts.accounting), (data.bank.file_name, data.bank.accounting)):
        rows.append([name, acc.total_rows, acc.preamble, acc.header, acc.noise, acc.parsed, acc.rejected,
                     "✔" if acc.balanced else "✘"])
    r = _table(ws, r, ["File", "Rows in sheet", "Above header", "Header", "Skipped (with reason)", "Parsed",
                       "Could not read", "Adds up"], rows)

    r += 1
    r = _section(ws, r, "Issues found")
    issues = [*data.payouts.issues, *data.bank.issues]
    rows = [[i.severity, i.source_file, i.source_row, i.code, i.message] for i in issues]
    r = _table(ws, r, ["Severity", "File", "Row", "Code", "Detail"], rows or [["—", "", None, "", "No issues"]],
               wrap_cols={5})

    r += 1
    r = _section(ws, r, f"Rows that could not be read (also listed on '{ACTION_SHEET}')")
    rows = [[rj.source_file, rj.source_row, rj.reason, "; ".join(f"{k}={v}" for k, v in rj.raw.items() if v)]
            for rj in (*data.payouts.rejected, *data.bank.rejected)]
    r = _table(ws, r, ["File", "Row", "Reason", "Raw values"], rows or [["—", None, "None", ""]], wrap_cols={3, 4})

    r += 1
    r = _section(ws, r, "Rows skipped on purpose")
    rows = [[n.source_file, n.source_row, n.reason, "; ".join(f"{k}={v}" for k, v in n.raw.items() if v)]
            for n in (*data.payouts.noise, *data.bank.noise)]
    _table(ws, r, ["File", "Row", "Reason", "Raw values"], rows or [["—", None, "None", ""]], wrap_cols={4})
    _widths(ws, [30, 14, 50, 70, 18, 10, 14, 10])


def _section(ws: Worksheet, r: int, title: str) -> int:
    ws.cell(r, 1, title).font = SECTION_FONT
    return r + 1


def _kv(ws: Worksheet, r: int, pairs: list[tuple[str, str]]) -> int:
    for k, v in pairs:
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return r


def _cell_value(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)  # Excel cells only hold doubles; the 2-decimal number format keeps the display exact
    return v


def _table(ws: Worksheet, r: int, headers: list[str], rows: Iterable[list], money_cols: set[int] = frozenset(),
           date_cols: set[int] = frozenset(), wrap_cols: set[int] = frozenset()) -> int:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    r += 1
    for row in rows:
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, _cell_value(v))
            if c in money_cols:
                cell.number_format = MONEY_FMT
            elif c in date_cols:
                cell.number_format = DATE_FMT
            if c in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
        r += 1
    return r


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _finish(ws: Worksheet, widths: list[int]) -> None:
    _widths(ws, widths)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
